"""An editable APA article scaffold; study claims remain the author's work."""

from io import BytesIO
import re


def manuscript_blocks(result: dict | None = None) -> list[tuple[str, str]]:
    result = result or {}
    config = result.get("config") or {}
    prep = result.get("prep") or {}

    def recorded(mapping, key, placeholder):
        value = mapping.get(key)
        return str(value) if value is not None and str(value).strip() else f"[{placeholder}]"

    method = recorded(config, "method", "estimation method")
    facets = ", ".join(map(str, config.get("facet_names") or [])) or "[facet names]"
    setup = (
        f"The analysis configuration specified {recorded(config, 'model', 'model')} "
        f"with {method} estimation in the standalone Python MFRM application "
        f"(version {recorded(config, 'app_version', 'version')}). "
        f"The retained data comprised {recorded(prep, 'n_obs', 'number of ratings')} ratings "
        f"and {recorded(prep, 'n_person', 'number of person identifiers')} person identifiers. "
        f"The selected facets were {facets}. "
        "[Identify whether these are empirical, secondary, or synthetic data; distinguish "
        "ratings from independent participants and describe the sampling units.]"
    )
    estimation = (
        "[Explain why this model answers the research question. State the facet signs, "
        "identification constraints, category parameterization, anchoring, estimator, "
        "optimizer, stopping tolerance, and uncertainty method actually used. "
        "Distinguish conditional standard errors from uncertainty that propagates "
        "parameter estimation. Cite the methods actually used.]"
    )
    if method.upper() == "MML":
        estimation += (
            " [For MML, report the quadrature rule and number of nodes, latent distribution, "
            "fixed or estimated population scale, latent regression if used, and numerical "
            "sensitivity. MML is not Bayesian posterior sampling. Check the current "
            "population-scale inference hold before reporting uncertainty.]"
        )
    elif method.upper() == "JMLE":
        estimation += " [For JMLE, report the treatment of extreme scores and any bias correction.]"
    return [
        ("title", "[Informative title identifying the construct population and study question]"),
        ("center", "[Author names]"),
        ("center", "[Department and institution for each author]"),
        ("heading", "Author Note"),
        ("body", "[ORCIDs, author affiliations, correspondence, funding, conflicts of interest, "
         "related reports using these data, and required acknowledgments. Follow the target "
         "journal's anonymization and disclosure instructions.]"),
        ("break", ""),
        ("heading", "Abstract"),
        ("abstract", "[Write this after the main text. Explain the practical problem and objective, "
         "the design and sample, the principal finding with its magnitude and uncertainty, "
         "and the answer with its scope. Do not replace the finding with a list of analyses. "
         "Follow the journal's word limit and structured or unstructured abstract format.]"),
        ("body", "Keywords: [construct], [assessment context], many-facet Rasch measurement"),
        ("break", ""),
        ("title", "[Repeat the manuscript title]"),
        ("body", "[Explain who encounters the assessment problem, in what setting, and why its "
         "consequences matter. Introduce the construct and a concrete rating example before "
         "technical terminology.]"),
        ("body", "[Synthesize the relevant evidence with citations. Explain what remains unresolved "
         "and what answering it would change. Separate the substantive gap from a software "
         "or computation improvement.]"),
        ("subheading", "The Present Study"),
        ("body", "[State RQ1 and any additional questions. For each, explain why the chosen data, "
         "comparisons, and outcomes can answer it. Distinguish prespecified hypotheses from "
         "exploratory questions. If using an empirical example and simulations, explain "
         "their distinct roles and connection to the central problem.]"),
        ("heading", "Method"),
        ("subheading", "Participants and Setting"),
        ("body", "[Report the population, sampling and recruitment, setting and dates, eligibility, "
         "exclusions, sample-size rationale, relevant characteristics, consent, and ethics "
         "approval or documented exemption. Do not invent an approval or preregistration.]"),
        ("subheading", "Materials and Rating Design"),
        ("body", "[Describe the tasks, rubric and category anchors, raters and training, allocation "
         "and masking, ratings per performance, shared ratings and design connectedness. "
         "Distinguish planned missing cells from omitted ratings and state the exclusion "
         "and missing-data rules. Explain why the design supports the intended comparisons.]"),
        ("subheading", "Analysis"),
        ("body", setup),
        ("body", estimation),
        ("body", "[For each research question, name the estimate or contrast, its direction and "
         "scale, uncertainty, and sensitivity analysis. State which fit, category, residual "
         "PCA, local-dependence, and bias checks were planned and which were computed. "
         "For bias, identify screened facet pairs, sparse-cell rules, practical thresholds, "
         "and the multiplicity correction and family. Explain decisions made after inspection.]"),
        ("heading", "Results"),
        ("subheading", "Data and Model Checks"),
        ("body", "[Report the retained sample, excluded or missing ratings, category support, "
         "convergence and estimation problems. Give the relevant fit and stability evidence. "
         "PCA is an exploratory residual-pattern check, not proof of unidimensionality. "
         "A nonsignificant bias check does not establish fairness or the absence of bias.]"),
        ("subheading", "[RQ1 Descriptive Heading]"),
        ("body", "[Answer RQ1 in words, then give the direction, size, unit, and uncertainty of "
         "the relevant estimate or contrast, with the table or figure number. Explain the "
         "practical meaning and applicable conditions. Repeat this structure for each RQ; "
         "retain caveats from the app's Results draft and evidence checks.]"),
        ("body", "[Statistical sentence to complete only when supported: The estimated contrast "
         "was [value] logits (SE = [value], 95% CI [[lower], [upper]]), using [contrast and "
         "uncertainty method]. If a valid test was computed, report its statistic, degrees "
         "of freedom, and exact p value or p < .001, and identify any adjusted p value. "
         "Never write p = .000 or infer a contrast SE from separate SEs without covariance.]"),
        ("subheading", "Sensitivity and Exploratory Findings"),
        ("body", "[Report analyses that were performed, including unfavorable results. Identify "
         "post hoc category changes or exclusions. Explain whether and how the answer "
         "changed. Simulation findings concern the stated generating conditions; they do "
         "not by themselves validate the empirical application.]"),
        ("heading", "Discussion"),
        ("body", "[Return to the original problem and answer each research question. Compare the "
         "finding with prior work, distinguishing evidence from possible explanations. "
         "Explain the implications for the intended assessment decision.]"),
        ("subheading", "Limitations and Scope"),
        ("body", "[Explain the consequences of sampling, rating overlap, estimation and uncertainty "
         "limits, model assumptions, and exploratory decisions. State what populations, "
         "designs, comparisons, or causal claims the evidence does not support. Use the "
         "current run's unresolved checks to make these limitations specific.]"),
        ("subheading", "Conclusion"),
        ("body", "[Give the bounded answer and why it matters. Do not introduce new results, "
         "equate convergence with validity, or declare the instrument valid from fit alone.]"),
        ("subheading", "Data and Code Availability"),
        ("body", "[Provide actual repository links or explain access restrictions. Report software "
         "versions, analysis settings, seeds where relevant, and which deidentified data "
         "and materials are available. Include preregistration and deviations if applicable; "
         "do not claim that materials are public before they are deposited.]"),
        ("break", ""),
        ("heading", "References"),
        ("reference", "[Insert only works cited in the completed manuscript, using Zotero's APA "
         "7th edition style. Check author names, year, title, journal, volume, pages, and DOI. "
         "These placeholders are ordinary text, not live Zotero citation fields.]"),
    ]


def manuscript_markdown(result: dict | None = None) -> str:
    prefixes = {"title": "# ", "heading": "## ", "subheading": "### "}
    return "\n\n".join(prefixes.get(kind, "") + text
                       for kind, text in manuscript_blocks(result) if kind != "break") + "\n"


def manuscript_word_bytes(result: dict | None = None) -> bytes:
    """APA professional-paper baseline, with editable placeholders and no raw rows."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Header"):
        style = doc.styles[name]
        style.font.name, style.font.size = "Times New Roman", Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + attr), None)
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        fmt = style.paragraph_format
        fmt.line_spacing = 2
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.widow_control = True
        fmt.first_line_indent = Inches(0.5) if name == "Normal" else Inches(0)
        if name in ("Title", "Heading 1", "Heading 2"):
            style.font.bold = True
            fmt.keep_with_next = True
            fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT if name == "Heading 2" else WD_ALIGN_PARAGRAPH.CENTER
    header = section.header.paragraphs[0]
    doc.styles["Header"].paragraph_format.tab_stops.clear_all()
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    header.add_run("[SHORT TITLE]\t")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    header._p.append(field)
    styles = {"title": "Title", "heading": "Heading 1", "subheading": "Heading 2"}
    for kind, text in manuscript_blocks(result):
        if kind == "break":
            doc.add_page_break()
            continue
        paragraph = doc.add_paragraph(style=styles.get(kind, "Normal"))
        for part in re.split(r"\b(SE|p)\b", text):
            run = paragraph.add_run(part)
            if part in {"SE", "p"}:
                run.italic = True
        if kind == "title" and len(doc.paragraphs) == 1:
            paragraph.paragraph_format.space_before = Pt(48)
        if kind in ("center", "abstract"):
            paragraph.paragraph_format.first_line_indent = Inches(0)
        if kind == "center":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if kind == "reference":
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
