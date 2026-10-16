"""Real input-to-export journeys; browser rendering and human understanding are separate."""
from io import BytesIO

import numpy as np
import pandas as pd
import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

import streamlit_app as app
from tests.test_help_contextual_app_integration import _block_analysis_calls, _locale_value
from tests.test_jmle_facet_boundary import separated_data


def _run(at):
    at.run(timeout=120)
    assert not at.exception, [error.value for error in at.exception]


def _prepare(case):
    at = AppTest.from_string("import streamlit_app as app\napp.main()", default_timeout=120)
    _run(at)
    at.button(key="onboarding_dismiss").click()
    _run(at)
    data = separated_data() if case == "facet_boundary" else app._make_self_test_gradient_data()
    at.text_area(key="paste_data_text").set_value(data.to_csv(index=False))
    _run(at)
    at.radio(key="facets_mode_workflow_mode").set_value("Advanced controls")
    _run(at)
    at.number_input(key="facets_mode_maxit").set_value(100)
    if case in {"fixed_sd", "free_sd"}:
        at.radio(key="facets_mode_estimation_method").set_value("MML")
        _run(at)
        at.selectbox(key="facets_mode_mml_engine").set_value("EM")
        at.checkbox(key="facets_mode_estimate_population_sd").set_value(case == "free_sd")
    at.radio(key="lang").set_value("ja")
    _run(at)
    assert next(w.value for w in at.selectbox if w.key and "person_col" in w.key) == "Person"
    assert next(w.value for w in at.selectbox if w.key and "score_col" in w.key) == "Score"
    return at


def _fit(at):
    label = _locale_value(at.session_state["lang"], "sidebar_perf.run_button")
    next(button for button in at.sidebar.button if button.label == label).click()
    _run(at)
    return at.session_state["facets_mode_output"]


@pytest.mark.parametrize("case", ["regular", "facet_boundary", "fixed_sd", "free_sd"])
def test_real_fit_help_home_language_exports_and_stale_input(case, monkeypatch):
    at = _prepare(case)
    output = _fit(at)
    result = output["result"]
    before = output["diagnostics"]["measures"].copy(deep=True)
    analysis_id = app.build_result_analysis_identity(result).analysis_id
    held = case in {"facet_boundary", "free_sd"}
    assert before.SE_Status.eq("calibration_inference_withheld").all() == held

    calls = _block_analysis_calls(monkeypatch)
    at.button_group(key="guided_essential_section").set_value("diagnostics")
    _run(at)
    assert at.session_state["guided_diagnostics_panel"] == "fit_details"
    at.button(key="mfrm_help_open_global").click()
    _run(at)
    at.radio(key="lang").set_value("en")
    _run(at)
    at.button(key="mfrm_help_return_to_source").click()
    _run(at)
    assert at.session_state["guided_essential_section"] == "diagnostics"
    at.button(key="mfrm_home_open").click()
    _run(at)
    at.button(key="mfrm_home_resume").click()
    _run(at)

    # Exercise actual document builders; figure rendering has separate coverage.
    monkeypatch.setattr(app, "_publication_figure_payloads", lambda *args: [])
    builds = []
    for kind in ("pdf", "word", "html"):
        name = f"build_publication_{kind}_bytes"
        original = getattr(app, name)

        def build(*args, _kind=kind, _original=original, **kwargs):
            builds.append(_kind)
            return _original(*args, **kwargs)

        monkeypatch.setattr(app, name, build)

    downloads = {}
    original_download = app.st.download_button

    def capture(label, data, *args, **kwargs):
        downloads[kwargs.get("file_name")] = data
        return original_download(label, data, *args, **kwargs)

    monkeypatch.setattr(app.st, "download_button", capture)
    at.button_group(key="guided_essential_section").set_value("report_export")
    _run(at)
    assert downloads["mfrm_publication_document.pdf"].startswith(b"%PDF")
    assert at.session_state["_publication_doc_pdf_fp"]
    _run(at)
    at.radio(key="publication_doc_format").set_value("word")
    _run(at)
    word = Document(BytesIO(downloads["mfrm_publication_document.docx"]))
    expected, _ = app._publication_measure_table(result, output["diagnostics"])
    rows = [[cell.text for cell in row.cells] for row in word.tables[0].rows]
    assert rows[0] == expected.columns.tolist()
    actual = pd.DataFrame(rows[1:], columns=rows[0])
    for column in expected:
        if pd.api.types.is_numeric_dtype(expected[column]):
            np.testing.assert_allclose(actual[column].astype(float), expected[column], atol=.0005)
        else:
            assert actual[column].tolist() == expected[column].astype(str).tolist()
    at.radio(key="publication_doc_format").set_value("html")
    _run(at)
    assert expected.to_html(index=False, border=0) in downloads["mfrm_publication_document.html"].decode()
    at.radio(key="publication_doc_format").set_value("pdf")
    _run(at)
    assert builds == ["pdf", "word", "html"]
    at.radio(key="lang").set_value("ja")
    _run(at)
    assert builds == ["pdf", "word", "html", "pdf"]
    if held:
        assert _locale_value("ja", "guided.export_state_hold") in [w.value for w in at.warning]

    at.button_group(key="guided_export_task").set_value("files")
    _run(at)
    at.selectbox(key="download_single_table").select("measures")
    _run(at)
    public_csv = pd.read_csv(BytesIO(downloads["mfrm_measures.csv"]))
    assert not public_csv.Facet.eq("Person").any()
    at.checkbox(key="downloads_public_export_mode").uncheck()
    _run(at)
    csv = pd.read_csv(BytesIO(downloads["mfrm_measures.csv"]))
    np.testing.assert_allclose(csv.Estimate, before.Estimate, equal_nan=True)
    np.testing.assert_allclose(csv.SE, before.SE, equal_nan=True)
    assert csv.SE_Status.tolist() == before.SE_Status.tolist()
    assert csv.CI_Status.tolist() == before.CI_Status.tolist()
    assert calls == {"estimate": 0, "refit": 0}
    assert at.session_state["facets_mode_output"] is output
    assert app.build_result_analysis_identity(result).analysis_id == analysis_id
    pd.testing.assert_frame_equal(output["diagnostics"]["measures"], before, check_exact=True)

    at.number_input(key="facets_mode_maxit").set_value(110)
    for lang in ("ja", "en"):
        at.radio(key="lang").set_value(lang)
        _run(at)
        assert _locale_value(lang, "app.stale_results_warning") in [w.value for w in at.warning]
        assert at.button(key="stale_rerun").label == _locale_value(lang, "app.run_updated_settings")
        result_download_keys = ("-dl_selected_csv", "-dl_tables_zip", "-dl_tables_xlsx",
                                "-dl_tables_html", "-publication_doc_pdf_dl")
        assert not [b.proto.id for b in at.get("download_button")
                    if b.proto.id.endswith(result_download_keys)]
        assert not any(w.key == "guided_essential_section" for w in at.button_group)
    assert calls == {"estimate": 0, "refit": 0}
    assert at.session_state["facets_mode_output"] is output
    at.number_input(key="facets_mode_maxit").set_value(100)
    _run(at)
    assert at.session_state["facets_mode_output"] is output
    assert not any(button.key == "stale_rerun" for button in at.button)
    assert calls == {"estimate": 0, "refit": 0}


@pytest.mark.parametrize("lang", ["ja", "en"])
def test_invalid_score_mapping_can_be_corrected_without_losing_settings(lang):
    at = _prepare("regular")
    at.radio(key="lang").set_value(lang)
    _run(at)
    score_key = next(w.key for w in at.selectbox if w.key and "score_col" in w.key)
    at.selectbox(key=score_key).select("Person")
    _run(at)
    assert at.error
    label = _locale_value(lang, "sidebar_perf.run_button")
    next(button for button in at.sidebar.button if button.label == label).click()
    _run(at)
    assert at.error and "facets_mode_output" not in at.session_state
    at.selectbox(key=score_key).select("Score")
    _run(at)
    output = _fit(at)
    assert output["score_col"] == "Score"
    assert output["result"]["prep"]["n_obs"] == 12
    assert at.number_input(key="facets_mode_maxit").value == 100
    assert at.radio(key="lang").value == lang


def test_restored_legacy_flags_stay_held_when_current_settings_differ(monkeypatch):
    at = _prepare("free_sd")
    # Opt in before the fit; emulate older saved uncertainty labels in its snapshot.
    at.session_state[app._RUN_HISTORY_STORE_SNAPSHOTS_KEY] = True
    output = _fit(at)
    entry = at.session_state[app._RUN_HISTORY_KEY][-1]
    saved = entry["output_snapshot"]["diagnostics"]["measures"]
    saved["SE_Status"] = "ok"
    saved["CI_Status"] = "ok"
    before = saved.copy(deep=True)
    identity = app.build_result_analysis_identity(output["result"]).analysis_id
    calls = _block_analysis_calls(monkeypatch)
    at.button_group(key="guided_essential_section").set_value("start")
    _run(at)
    at.button(key=f"restore_{entry['run_id']}").click()
    _run(at)
    restored = at.session_state["facets_mode_output"]
    assert restored is not output
    assert app.build_result_analysis_identity(restored["result"]).analysis_id == identity
    at.number_input(key="facets_mode_maxit").set_value(110)
    for lang in ("ja", "en"):
        at.radio(key="lang").set_value(lang)
        _run(at)
        message = _locale_value(lang, "app.restored_results_info").format(timestamp=entry["timestamp"])
        assert message in [w.value for w in at.info]
        assert at.button(key="restored_run_current_settings").label == _locale_value(lang, "app.run_current_settings")
    monkeypatch.setattr(app, "_publication_figure_payloads", lambda *args: [])
    at.button_group(key="guided_essential_section").set_value("report_export")
    _run(at)
    at.radio(key="publication_doc_format").set_value("word")
    _run(at)
    word = Document(BytesIO(at.session_state["_publication_doc_word_bytes"]))
    columns = [cell.text for cell in word.tables[0].rows[0].cells]
    for row in word.tables[0].rows[1:]:
        values = dict(zip(columns, [cell.text for cell in row.cells]))
        assert values["Basis"] == "Withheld"
        assert values["SE / SD"] == values["Lower"] == values["Upper"] == ""
    pd.testing.assert_frame_equal(saved, before, check_exact=True)
    assert at.session_state["facets_mode_output"] is restored
    assert calls == {"estimate": 0, "refit": 0}
