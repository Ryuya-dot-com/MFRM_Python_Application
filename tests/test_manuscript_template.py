"""Keep author judgments separate from run facts and preserve the writing route."""

from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from mfrm_app.manuscript import manuscript_markdown, manuscript_word_bytes


def test_manuscript_uses_recorded_facts_without_inventing_results_or_exposing_rows():
    result = {"config": {"model": "PCM", "method": "MML", "app_version": "test",
                          "facet_names": ["Rater", "Task"], "estimate_population_sd": True},
              "prep": {"n_obs": 120, "n_person": 10},
              "raw_data": "SECRET_PERSON_NAME"}
    text = manuscript_markdown(result)
    assert "120 ratings and 10 person identifiers" in text
    assert "PCM with MML estimation" in text
    assert "quadrature" in text and "inference hold" in text
    assert "[Answer RQ1" in text and "[Give the bounded answer" in text
    assert "SECRET_PERSON_NAME" not in text
    assert "## References" in text and "not live Zotero citation fields" in text
    assert "[number of ratings]" in manuscript_markdown()
    changed = {**result, "prep": {"n_obs": 121, "n_person": 11}}
    assert "121 ratings and 11 person identifiers" in manuscript_markdown(changed)
    assert result["prep"]["n_obs"] == 120


def test_word_scaffold_has_apa_layout_and_same_content_as_markdown():
    data = manuscript_word_bytes()
    doc = Document(BytesIO(data))
    assert doc.sections[0].left_margin.inches == 1
    assert doc.sections[0].page_width.inches == 8.5
    assert doc.styles["Normal"].font.name == "Times New Roman"
    assert doc.styles["Normal"].font.size.pt == 12
    assert doc.styles["Normal"].paragraph_format.line_spacing == 2
    assert doc.styles["Normal"].paragraph_format.first_line_indent.inches == 0.5
    assert doc.paragraphs[0].style.name == "Title"
    assert not doc.styles["Title"].element.xpath("./w:pPr/w:pBdr")
    assert not doc.styles["Heading 1"].element.xpath("./w:rPr/w:rFonts/@w:asciiTheme")
    for heading in ("Abstract", "Method", "Results", "Discussion", "References"):
        assert any(p.text == heading and p.style.name == "Heading 1" for p in doc.paragraphs)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "[Answer RQ1" in text and "[number of ratings]" in text
    assert doc.paragraphs[-1].paragraph_format.first_line_indent.inches == -0.5
    with ZipFile(BytesIO(data)) as zipped:
        assert b'PAGE' in zipped.read("word/header1.xml")


@pytest.mark.parametrize("lang", ["en", "ja"])
def test_manuscript_ui_keeps_result_generation_opt_in_and_current_context(lang):
    def view(lang):
        from unittest.mock import patch
        import pandas as pd
        import streamlit as st
        import streamlit_app as app
        st.session_state["lang"] = lang
        st.session_state.setdefault("fit", {"config": {"model": "RSM", "method": "MML",
                                                       "estimate_population_sd": True}})
        st.session_state.setdefault("draft_calls", 0)
        context = pd.DataFrame({"StopBeforeFinalOutput": [True]})
        def draft(*args, **kwargs):
            assert kwargs["simulation_sparse_context"].equals(context)
            assert kwargs["bias_results"] == {"facet_a": "Rater", "facet_b": "Task"}
            st.session_state["draft_calls"] += 1
            return "Editing worksheet only. Keep the inference hold."
        with patch.object(app, "generate_report_ready_apa_results_draft", draft), \
             patch.object(app, "current_custom_simulation_sparse_export_frames",
                          return_value={"custom_simulation_sparse_reporting_context": context}):
            app._render_manuscript_template_section(st.session_state["fit"], {},
                bias_results={"facet_a": "Rater", "facet_b": "Task"})

    at = AppTest.from_function(view, args=(lang,)).run(timeout=45)
    assert not at.exception and at.session_state["draft_calls"] == 0
    assert len(at.get("download_button")) == 1
    assert all(not e.proto.expanded for e in at.expander)
    assert not any("guided.manuscript" in c.value for c in at.caption)
    at.radio(key="manuscript_format").set_value("markdown").run()
    assert not at.exception and at.session_state["draft_calls"] == 0
    at.checkbox(key="manuscript_results_preview").check().run()
    assert not at.exception and at.session_state["draft_calls"] == 1
    assert any("Keep the inference hold" in m.value for m in at.markdown)
    assert at.session_state["fit"]["config"]["estimate_population_sd"] is True
