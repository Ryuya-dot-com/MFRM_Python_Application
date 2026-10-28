"""Real fits retain uncertainty meaning across tables and publication formats."""
from io import BytesIO, StringIO

import numpy as np
import pandas as pd
import pytest
from docx import Document

import streamlit_app as app
from tests.test_jmle_facet_boundary import separated_data


@pytest.fixture(scope="module", params=["regular", "facet_boundary", "fixed_sd", "free_sd"])
def fitted_case(request):
    case = request.param
    data = separated_data() if case == "facet_boundary" else app._make_self_test_gradient_data()
    options = ({"method": "MML", "mml_engine": "EM", "quad_points": 15,
                "estimate_population_sd": case == "free_sd"}
               if case in {"fixed_sd", "free_sd"} else {"method": "JMLE"})
    result = app.mfrm_estimate(data, "Person", ["Rater", "Task"], "Score", maxit=100, **options)
    diagnostics = app.mfrm_diagnostics(result, compute_pca=False, interaction_pairs=[])
    return case, result, diagnostics


def test_real_measures_preserve_values_and_current_holds_in_csv_and_scaled_tables(fitted_case):
    case, result, diagnostics = fitted_case
    held = case in {"facet_boundary", "fixed_sd", "free_sd"}
    original = diagnostics["measures"].copy(deep=True)
    frames = app.build_result_bundle_frames(result, diagnostics)
    downloads, _ = app.collect_download_frames(result, diagnostics, {}, pd.DataFrame(), pd.DataFrame(),
                                              public_export_mode=False)
    apa = app._collect_apa_exportable_tables(result, diagnostics)["Combined measures and fit"]
    for table in (frames["measures"], downloads["measures"], apa):
        csv = pd.read_csv(StringIO(table.to_csv(index=False)))
        np.testing.assert_allclose(csv.Estimate, original.Estimate, equal_nan=True)
        np.testing.assert_allclose(csv.SE, original.SE, equal_nan=True)
        assert csv.SE_Status.eq("calibration_inference_withheld").all() == held
        if held:
            assert csv.CI_Status.eq("calibration_inference_withheld").all()
            assert csv.UncertaintyCaution.str.contains("inference is withheld").all()
    scaled = app.calc_facets_report_tbls(result, diagnostics, uscale=2.)
    for facet, table in scaled.items():
        source = original[original.Facet.eq(facet)].set_index("Level").reindex(table.Level)
        np.testing.assert_allclose(table.ModelSE, 2*source.SE, atol=.005, equal_nan=True)
        assert table.SE_Status.tolist() == source.SE_Status.tolist()
        assert table.SE_Method.tolist() == source.SE_Method.tolist()
    pd.testing.assert_frame_equal(diagnostics["measures"], original)


def test_writing_appendix_uses_current_scale_and_uncertainty_qualifications(fitted_case):
    case, result, diagnostics = fitted_case
    text = app.generate_manuscript_reporting_template(result, diagnostics)
    assert "Do not imply that the latent variance was freely estimated" not in text
    assert "Limit no-bias statements" not in text
    assert "Bias/no-bias conclusions remain withheld" in text
    assert "calibration_inference_withheld:" not in text
    assert "conditional_approximation:" not in text
    assert "publication gate status" not in text
    assert "**Overall manuscript gate**" not in text
    assert "A development calculation" not in text
    if case in {"facet_boundary", "fixed_sd", "free_sd"}:
        assert "interpretation is on hold" in text
        guide = app.build_manuscript_claim_guide(result, diagnostics)
        for area in ("Reliability and separation", "Wright map targeting and measure interpretation",
                     "Rating-scale functioning", "Fit and dimensionality", "Anchoring and linking",
                     "Prediction, plausible values, and simulation"):
            row = guide.loc[guide.ManuscriptArea.eq(area)].iloc[0]
            assert row.ClaimStatus == "Do not claim"
            assert "Interpretation is withheld" in row.SafeManuscriptWording
            assert row.EvidenceToReport.startswith("Technical checks only")
        assert "High rater reliability was interpreted as evidence" not in text
        assert "Final parameter interpretation was withheld" in text
        assert "STAT-003" not in text and "G2" not in text
    elif case == "fixed_sd":
        assert "not the sampling SE of EAP" in text
    else:
        assert "other fitted parameters are held fixed" in text
    scale = next(line for line in text.splitlines() if line.startswith("- MML population scale:"))
    if case == "free_sd":
        assert "estimated" in scale.lower() and "withheld" in scale.lower()
    elif case == "fixed_sd":
        assert "fixed" in scale.lower()
    else:
        assert "not MML" in scale


def test_real_publication_formats_keep_basis_and_suppress_held_intervals(fitted_case, monkeypatch):
    case, result, diagnostics = fitted_case
    held = case in {"facet_boundary", "fixed_sd", "free_sd"}
    # Figure rendering has its own tests; exercise the actual document builders and tables.
    monkeypatch.setattr(app, "_publication_figure_payloads", lambda *args: [])
    table, caption = app._publication_measure_table(result, diagnostics)
    assert len(table.columns) <= 10
    assert {"Infit", "Outfit"}.issubset(table.columns)
    measures = diagnostics["measures"]
    person_held = measures.Facet.eq("Person") & measures.get(
        "PersonInferenceReady", pd.Series(True, index=measures.index)).eq(False).fillna(False)
    for name in ("Infit", "Outfit"):
        np.testing.assert_allclose(pd.to_numeric(table[name], errors="coerce"),
                                   measures[name].mask(person_held), atol=.0005, equal_nan=True)
    assert "not automatic pass/fail" in caption
    assert table.Basis.eq("Withheld").all() == held
    if held:
        assert table[["SE / SD", "Lower", "Upper"]].eq("").all().all()
        assert "technical terminal-point" in caption
    elif case == "fixed_sd":
        assert table.loc[table.Facet.eq("Person"), "Basis"].eq("Posterior SD").all()
        assert "not the sampling SE of EAP" in caption
    else:
        assert table.Basis.eq("Conditional").all()
    basic_title, basic_table, basic_note = next(item for item in app._publication_report_tables(
        result, diagnostics, {}, include_details=False) if item[0] == "Facet Estimates and Fit")
    assert basic_title == "Facet Estimates and Fit"
    assert not basic_table.Facet.eq("Person").any()
    np.testing.assert_allclose(pd.to_numeric(basic_table.Estimate),
                               diagnostics["measures"].loc[diagnostics["measures"].Facet.ne("Person"), "Estimate"],
                               atol=.0005)
    if held:
        assert "Basis = Withheld" in basic_note
        assert basic_table["Interval"].eq("").all()
    if case == "fixed_sd":
        assert "Posterior SD" not in basic_note
    reliability, reliability_caption = app._publication_reliability_table(result, diagnostics)
    assert reliability.Use.eq("Withheld").all() == held
    if held:
        assert reliability[["RMSE", "Separation", "Strata", "Reliability"]].eq("").all().all()
        assert "Interpretation is withheld" in reliability_caption
    raw_basis = table.Basis.copy()
    report_tables = app._publication_report_tables(result, diagnostics, {}, include_details=True)
    measure_idx = next(i for i, item in enumerate(report_tables) if item[0] == "Element Estimates and Fit")
    reliability_idx = next(i for i, item in enumerate(report_tables) if item[0] == "Reliability and Separation by Facet")
    _, table, caption = report_tables[measure_idx]
    _, reliability, reliability_caption = report_tables[reliability_idx]
    word = Document(BytesIO(app.build_publication_word_bytes(result, diagnostics)))
    assert [cell.text for cell in word.tables[measure_idx].rows[0].cells] == table.columns.tolist()
    if "Basis" in table:
        basis_col = table.columns.get_loc("Basis")
        assert [row.cells[basis_col].text for row in word.tables[measure_idx].rows[1:]] == raw_basis.tolist()
    else:
        assert f"Basis = {raw_basis.iloc[0]}" in caption
    if held:
        assert table["Interval"].eq("").all()
    assert any(caption in paragraph.text for paragraph in word.paragraphs)
    assert any(reliability_caption in paragraph.text for paragraph in word.paragraphs)
    assert [cell.text for cell in word.tables[reliability_idx].rows[0].cells] == reliability.columns.tolist()
    html = app.build_publication_html_bytes(result, diagnostics).decode()
    assert app._html_escape(caption) in html
    html_tables = pd.read_html(StringIO(html))
    assert html_tables[measure_idx].columns.tolist() == table.columns.tolist()
    np.testing.assert_allclose(html_tables[measure_idx].Estimate, pd.to_numeric(table.Estimate), atol=.0005)
    assert html_tables[reliability_idx].columns.tolist() == reliability.columns.tolist()
    if held:
        assert html_tables[measure_idx]["Interval"].isna().all()
    assert "<em>Note.</em>" in html
    pdf_tables = []
    original_pdf_table = app._pdf_dataframe_flowable
    def record_table(frame, caption=None, **kwargs):
        pdf_tables.append((frame, kwargs["note"]))
        return original_pdf_table(frame, caption, **kwargs)
    monkeypatch.setattr(app, "_pdf_dataframe_flowable", record_table)
    pdf = app.build_publication_pdf_bytes(result, diagnostics)
    assert pdf.startswith(b"%PDF")
    pd.testing.assert_frame_equal(pdf_tables[measure_idx][0], table)
    pd.testing.assert_frame_equal(pdf_tables[reliability_idx][0], reliability)
    assert caption in pdf_tables[measure_idx][1]


def test_category_and_step_exports_match_the_diagnostic_tables(fitted_case):
    _, result, diagnostics = fitted_case
    before = diagnostics["obs"].copy(deep=True)
    frames, _ = app.collect_download_frames(result, diagnostics, {}, pd.DataFrame(), pd.DataFrame(),
                                            public_export_mode=True)
    category = app.calc_category_stats(diagnostics["obs"], result)
    steps = app.calc_step_order(result["steps"])
    pd.testing.assert_frame_equal(frames["category_diagnostics"], category)
    # Exports may attach current inference holds to a saved step table.
    pd.testing.assert_frame_equal(frames["step_order"][steps.columns], steps)
    assert np.isclose(category.Count.sum(), diagnostics["obs"].Weight.sum())
    assert np.isclose(category.Percent.sum(), 100.)
    pd.testing.assert_frame_equal(diagnostics["obs"], before)


@pytest.mark.parametrize("config", [
    {"method": "MML", "estimate_population_sd": True},
    {"method": "JMLE", "jmle_facet_boundary_review": True},
])
def test_saved_positive_flags_do_not_override_current_hold_or_modify_source(config):
    source = pd.DataFrame({"Facet": ["Person"], "Level": ["P1"], "Estimate": [0.], "SE": [.2],
                           "CI_Lower": [-.392], "CI_Upper": [.392], "SE_Status": ["ok"], "CI_Status": ["ok"]})
    before = source.copy(deep=True)
    qualified = app.qualify_measure_uncertainty(source, {"config": config})
    assert qualified.SE_Status.eq("calibration_inference_withheld").all()
    assert qualified.CI_Status.eq("calibration_inference_withheld").all()
    assert qualified.SE.iloc[0] == .2
    table, _ = app._publication_measure_table({"config": config}, {"measures": source})
    assert table.Basis.iloc[0] == "Withheld"
    assert table["SE / SD"].iloc[0] == ""
    pd.testing.assert_frame_equal(source, before)


@pytest.mark.parametrize("lang", ["en", "ja"])
@pytest.mark.parametrize("case", ["facet_boundary", "free_sd"])
def test_actual_report_table_displays_uncertainty_statuses_and_hold(lang, case):
    from streamlit.testing.v1 import AppTest

    def view(lang, case):
        from contextlib import ExitStack
        from unittest.mock import patch
        import streamlit as st
        import streamlit_app as app
        from tests.test_jmle_facet_boundary import separated_data

        st.session_state["lang"] = lang
        data = separated_data() if case == "facet_boundary" else app._make_self_test_gradient_data()
        options = ({"method": "MML", "mml_engine": "EM", "quad_points": 15,
                    "estimate_population_sd": True} if case == "free_sd" else {"method": "JMLE"})
        result = app.mfrm_estimate(data, "Person", ["Rater", "Task"], "Score", maxit=100, **options)
        diagnostics = app.mfrm_diagnostics(result, compute_pca=False, interaction_pairs=[])
        # Optional planning and export widgets are independent of the report table.
        with ExitStack() as stack:
            for name in ("render_estimand_contract_panel", "render_model_choice_guidance",
                         "render_g_d_study_section", "render_design_network_section",
                         "render_rater_severity_section", "render_rater_halo_section",
                         "render_parameter_recovery_simulation", "render_apa_presets_section"):
                stack.enter_context(patch.object(app, name, return_value=None))
            app._render_report_tables(result, diagnostics)
        st.session_state["hold_message"] = app.t("guided.export_state_hold")

    at = AppTest.from_function(view, args=(lang, case)).run(timeout=60)
    assert not at.exception
    assert at.session_state["hold_message"] in [warning.value for warning in at.warning]
    table = next(frame.value for frame in at.dataframe if "CI_Status" in frame.value.columns)
    assert {"SE_Method", "SE_Status", "CI_Method", "UncertaintyCaution"}.issubset(table.columns)
    assert table.SE_Status.eq("calibration_inference_withheld").all()
    assert table.CI_Status.eq("calibration_inference_withheld").all()
