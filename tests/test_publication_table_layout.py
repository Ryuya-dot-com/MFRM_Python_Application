"""Reporting scope and APA structure retain values, qualifications and usable notes."""
from io import BytesIO, StringIO

import numpy as np
import pandas as pd
import pytest
from docx import Document
from docx.oxml.ns import qn
from reportlab.platypus import Paragraph, Table

import streamlit_app as app


@pytest.fixture
def report_data():
    result = {"config": {"model": "RSM", "method": "JMLE", "facet_names": ["Rater"]}}
    diagnostics = {"measures": pd.DataFrame({
        "Facet": ["Person", "Rater", "Rater"], "Level": ["P1", "R1", "R2"],
        "Estimate": [.8, -.2, .2], "SE": [.1, .2, .2],
        "Infit": [.9, 1.0, 1.1], "Outfit": [1.05, .95, 1.0],
    })}
    return result, diagnostics


def test_basic_report_keeps_results_first_and_details_are_opt_in(report_data, monkeypatch):
    result, diagnostics = report_data
    original = diagnostics['measures'].copy(deep=True)
    monkeypatch.setattr(app, "_publication_figure_payloads", lambda *args: [])
    calls = []
    monkeypatch.setattr(app, "generate_method_appendix_text", lambda *args: calls.append("methods") or "# Appendix method")
    monkeypatch.setattr(app, "generate_manuscript_reporting_template", lambda *args: calls.append("writing") or "# Writing example")
    for builder in (app.build_publication_html_bytes, app.build_publication_word_bytes, app.build_publication_pdf_bytes):
        assert builder(result, diagnostics, include_details=False)
    assert calls == []
    basic = app.build_publication_html_bytes(result, diagnostics, include_details=False).decode()
    assert "Appendix method" not in basic
    basic_table = pd.read_html(StringIO(basic))[0]
    assert basic_table.Level.tolist() == ["R1", "R2"]
    assert "Individual person estimates and fit: detailed report" in basic
    full = app.build_publication_html_bytes(result, diagnostics, include_details=True).decode()
    assert calls == ["methods", "writing"]
    assert full.index("<table") < full.index("Appendix method")
    assert pd.read_html(StringIO(full))[0].Level.tolist() == ["P1", "R1", "R2"]
    pd.testing.assert_frame_equal(diagnostics['measures'], original)


def test_basic_report_does_not_render_unused_figures(report_data, monkeypatch):
    result, diagnostics = report_data
    monkeypatch.setattr(app, "_plotly_or_matplotlib_png", lambda *args: b"png")
    def forbidden(*args, **kwargs):
        pytest.fail("An optional figure was rendered for the basic report")
    monkeypatch.setattr(app, "_make_fit_scatter_export_figure", forbidden)
    monkeypatch.setattr(app, "_make_facet_distribution_export_figure", forbidden)
    figures = app._publication_report_figures(result, diagnostics, include_details=False)
    assert len(figures) == 2
    assert figures[0][0] == "wright_map"
    assert figures[1][0].startswith("category_curves_")


def test_word_table_has_separate_number_title_and_note_without_grid():
    doc = Document()
    frame = pd.DataFrame({"Rater": ["A&B"], "Estimate": [0.125]})
    app._add_dataframe_to_docx(doc, frame, "Rater Estimates", table_number=2, note="Estimates are in logits.")
    assert [p.text for p in doc.paragraphs] == ["Table 2", "Rater Estimates", "Note. Estimates are in logits."]
    assert doc.paragraphs[0].runs[0].bold
    assert doc.paragraphs[1].runs[0].italic
    assert doc.paragraphs[2].runs[0].italic
    assert not doc.paragraphs[2].runs[1].italic
    table = doc.tables[0]
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    for edge in ("top", "bottom"):
        assert borders.find(qn(f"w:{edge}")).get(qn("w:val")) == "single"
    for edge in ("left", "right", "insideH", "insideV"):
        assert borders.find(qn(f"w:{edge}")).get(qn("w:val")) == "nil"
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert table.cell(1, 1).text == "0.125"
    assert not table._tbl.xpath('.//w:shd')
    assert table.cell(1, 0).paragraphs[0].paragraph_format.keep_with_next


def test_pdf_table_uses_horizontal_rules_and_separate_note():
    frame = pd.DataFrame({"Rater": ["A&B"], "Estimate": [0.125]})
    flowables = app._pdf_dataframe_flowable(frame, "Rater Estimates", table_number=2, note="SE < 1 & qualified.")
    table = next(f for f in flowables if isinstance(f, Table))
    text = [f.getPlainText() for f in flowables if isinstance(f, Paragraph)]
    assert text == ["Table 2", "Rater Estimates", "Note. SE < 1 & qualified."]
    assert [cmd[0] for cmd in table._linecmds] == ["LINEABOVE", "LINEBELOW", "LINEBELOW"]
    assert not table._bkgrndcmds
    assert table.repeatRows == 1
    assert table._cellvalues[1][1].getPlainText() == "0.125"


def test_all_formats_use_the_same_explicit_row_limit_and_keep_columns():
    frame = pd.DataFrame({f"C{i}": np.arange(70) for i in range(13)})
    display, note = app._publication_table_display(frame, "Units: logits.")
    assert display.shape == (50, 13)
    assert "first 50 of 70 rows" in note
    assert "complete table is in the CSV download" in note
    assert frame.shape == (70, 13)


def test_mixed_uncertainty_table_does_not_wrap_numeric_cells():
    frame = pd.DataFrame({"Facet": ["Person", "Rater"], "Level": ["P1", "R1"],
                          "Estimate": [-1.084, .2], "Infit": [.792, 1.025], "Outfit": [1.141, 1.025],
                          "SE / SD": [.553, .388], "Basis": ["Posterior SD", "Technical"],
                          "Interval": ["[-2.168, 0.000]", "[-0.561, 0.961]"]})
    table = next(f for f in app._pdf_dataframe_flowable(frame) if isinstance(f, Table))
    for j, column in enumerate(frame.columns):
        if pd.api.types.is_numeric_dtype(frame[column]):
            for i in range(1, len(frame) + 1):
                cell = table._cellvalues[i][j]
                style = table._cellStyles[i][j]
                width = table._colWidths[j] - style.leftPadding - style.rightPadding
                cell.wrap(width, 1000)
                assert len(cell.blPara.lines) == 1, (column, cell.getPlainText(), width)
