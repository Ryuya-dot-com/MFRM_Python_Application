"""Appendix formatting preserves qualifications, literal data and list hierarchy."""
from io import BytesIO
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml.ns import qn
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table

import streamlit_app as app
from mfrm_app import report_markup


SAMPLE = r"""# **Interpretation** <review>

- **Held result**: SE < 1 & not validated.
  - Read `SE_Status`; do not report *p* = .000.
  - Literal <img src="x"> and |ZSTD| >= 2.
- Next action: review the data.

| Field | Value |
| :--- | ---: |
| **Status** | Withheld |
| Literal | A\|B & <tag> |
| --- |  |

Do not turn a blank interval into zero.
"""


def test_shared_blocks_keep_data_rows_and_literal_operators():
    blocks = list(report_markup.blocks(SAMPLE))
    assert [depth for kind, depth, _ in blocks if kind == "bullet"] == [0, 1, 1, 0]
    table = next(content for kind, _, content in blocks if kind == "table")
    assert table[-1] == ["---", ""]
    assert table[2][1] == r"A\|B & <tag>"
    assert list(report_markup.inline_spans(r"A\|B *p* **Hold** `a_b`")) == [
        ("text", "A"), ("text", "|"), ("text", "B "), ("italic", "p"),
        ("text", " "), ("bold", "Hold"), ("text", " "), ("code", "a_b"),
    ]
    assert list(report_markup.blocks("|ZSTD| >= 2")) == [("paragraph", 0, "|ZSTD| >= 2")]


def test_word_appendix_has_real_emphasis_lists_and_minimal_table_borders():
    doc = Document()
    app._add_markdown_to_docx(doc, SAMPLE)
    assert doc.paragraphs[0].text == "Interpretation <review>"
    assert doc.paragraphs[0].runs[0].bold
    bullets = [p for p in doc.paragraphs if p.style.name.startswith("List Bullet")]
    assert [p.style.name for p in bullets] == ["List Bullet", "List Bullet 2", "List Bullet 2", "List Bullet"]
    assert bullets[1].paragraph_format.left_indent > bullets[0].paragraph_format.left_indent
    assert any(run.text == "SE_Status" and run.font.name == "Courier New" for run in bullets[1].runs)
    assert any(run.text == "p" and run.italic for run in bullets[1].runs)
    assert doc.tables[0].cell(2, 1).text == "A|B & <tag>"
    assert doc.tables[0].cell(1, 0).paragraphs[0].runs[0].bold
    borders = doc.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders.find(qn("w:insideV")).get(qn("w:val")) == "nil"
    assert doc.tables[0].rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None


def test_pdf_appendix_preserves_text_and_indents_wrapped_bullets():
    parts = app._pdf_markdown_to_flowables(SAMPLE, getSampleStyleSheet())
    paragraphs = [part for part in parts if isinstance(part, Paragraph)]
    assert paragraphs[0].getPlainText() == "Interpretation <review>"
    assert paragraphs[0].getKeepWithNext()
    bullets = [part for part in paragraphs if part.bulletText]
    assert [part.style.leftIndent for part in bullets] == [18, 36, 36, 18]
    assert "SE_Status" in bullets[1].getPlainText()
    assert "<img" in bullets[2].getPlainText()
    assert "**" not in "".join(p.getPlainText() for p in paragraphs)
    table = next(part for part in parts if isinstance(part, Table))
    assert table._cellvalues[2][1].getPlainText() == "A|B & <tag>"
    assert table.repeatRows == 1
    output = BytesIO()
    SimpleDocTemplate(output).build(parts)
    assert output.getvalue().startswith(b"%PDF")


def test_parent_bullet_stays_with_its_explanation():
    markdown = "- **Interpretation on hold**:\n  - Check the data before reporting."
    doc = Document()
    app._add_markdown_to_docx(doc, markdown)
    assert doc.paragraphs[0].paragraph_format.keep_with_next
    parts = app._pdf_markdown_to_flowables(markdown, getSampleStyleSheet())
    assert parts[0].getKeepWithNext()
    assert not parts[1].getKeepWithNext()


def test_html_appendix_has_nested_lists_and_escapes_raw_tags():
    html = app._markdown_to_html(SAMPLE)
    root = ET.fromstring(f"<main>{html}</main>")
    assert len(root.findall("./ul/li/ul/li")) == 2
    assert root.find(".//code").text == "SE_Status"
    assert root.find(".//thead/tr/th").text == "Field"
    assert root.find(".//img") is None
    text = "".join(root.itertext())
    assert "A|B & <tag>" in text and "|ZSTD| >= 2" in text
    assert "**" not in text and "`" not in text
    assert "blank interval into zero" in text


def test_appendix_tables_keep_all_rows_and_have_no_unrelated_csv_truncation_note():
    markdown = "| Field | Value |\n| --- | --- |\n" + "\n".join(f"| row{i} | {i} |" for i in range(55))
    doc = Document()
    app._add_markdown_to_docx(doc, markdown)
    assert len(doc.tables[0].rows) == 56
    assert doc.tables[0].cell(55, 1).text == "54"
    parts = app._pdf_markdown_to_flowables(markdown, getSampleStyleSheet())
    table = next(part for part in parts if isinstance(part, Table))
    assert len(table._cellvalues) == 56
    root = ET.fromstring(f"<main>{app._markdown_to_html(markdown)}</main>")
    assert len(root.findall(".//tbody/tr")) == 55
    assert not doc.paragraphs
