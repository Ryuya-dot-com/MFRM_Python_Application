"""Applied writing outputs follow retained data and current inference holds."""
from copy import deepcopy
from io import BytesIO, StringIO
from types import SimpleNamespace
from itertools import combinations

import numpy as np
import pandas as pd
import pytest
from docx import Document

import streamlit_app as app
from mfrm_app.manuscript import manuscript_markdown, manuscript_word_bytes
from mfrm_app.reporting_facts import observed_report_tables, optimizer_termination_text


@pytest.fixture
def writing_run():
    data = pd.DataFrame({"Person": ["P1"]*3 + ["P2"]*3,
                         "Rater": ["R1", "R2", "R1", "R2", "R1", "R2"],
                         "Score": [0, 0, 2, 2, 2, 3], "Weight": [1., 2., 1., 1., 2., 1.]})
    result = {"config": {"model": "PCM", "method": "JMLE", "facet_names": ["Rater"],
                          "facet_signs": {"Rater": 1}, "noncenter_facet": "Person",
                          "optimizer": "L-BFGS-B", "maxit": 200, "reltol": 1e-7,
                          "n_cat": 4, "facet_specs": {"Rater": {"centered": True}}},
              "prep": {"data": data, "n_obs": 999, "n_person": 99, "rating_min": 0, "rating_max": 3},
              "opt": SimpleNamespace(success=True, nit=12, message="stopping rule reached")}
    measures = pd.DataFrame({"Facet": ["Rater", "Rater"], "Level": ["R1", "R2"],
                             "Estimate": [-.4, .4], "SE": [.2, .2], "Infit": [.9, 1.1],
                             "Outfit": [.8, 1.2], "SE_Status": ["conditional_approximation"]*2,
                             "CI_Status": ["conditional_approximation"]*2})
    return result, {"measures": measures}


def test_observed_tables_separate_ratings_people_weights_and_unused_categories(writing_run):
    result, _ = writing_run
    original = result["prep"]["data"].copy(deep=True)
    design, categories = observed_report_tables(result)
    facts = design[1].set_index("Characteristic")["Value"]
    assert facts["Retained ratings"] == "6" and facts["Retained person identifiers"] == "2"
    assert facts["Rater levels"] == "2" and facts["Ratings per Rater level"] == "3–3"
    assert facts["Unique Rater levels per person"] == "2–2"
    assert not any(value in design[1].to_string() for value in ("P1", "P2", "R1", "R2"))
    table = categories[1]
    np.testing.assert_array_equal(table.Category, [0, 1, 2, 3])
    np.testing.assert_array_equal(table.Count, [2, 0, 3, 1])
    np.testing.assert_allclose(table["Weight sum"], [3, 0, 4, 1])
    np.testing.assert_allclose(table.Percent, [33.33, 0, 50, 16.67])
    assert "unweighted" in categories[2] and "separate likelihood weighting" in categories[2]
    pd.testing.assert_frame_equal(result["prep"]["data"], original)


@pytest.mark.parametrize("bounds", [(None, None), (np.nan, 3), (0, np.inf), (3, 0), (.5, 3), (0, 10**12)])
def test_malformed_saved_score_support_does_not_drop_scores_or_allocate_large_ranges(writing_run, bounds):
    result, _ = writing_run
    result["prep"].update(rating_min=bounds[0], rating_max=bounds[1])
    _, (_, categories, note) = observed_report_tables(result)
    assert len(categories) == 3 and categories.Count.sum() == 6
    assert categories.attrs["ReasonCode"] == "report.score_support_unavailable"
    assert "only observed categories are listed" in note


def test_observed_score_outside_saved_support_is_retained(writing_run):
    result, _ = writing_run
    result["prep"]["rating_max"] = 2
    _, (_, categories, note) = observed_report_tables(result)
    assert categories.Count.sum() == 6 and 3 in categories.Category.tolist()
    assert categories.attrs["ReasonCode"] == "report.observed_outside_recorded_support"
    assert "outside the recorded support" in note


def test_nullable_missing_values_cannot_silently_reduce_category_totals(writing_run):
    result, _ = writing_run
    result["prep"]["data"]["Weight"] = pd.Series([2, pd.NA, 1, 1, 1, 1], dtype="Float64")
    assert "Weight sum" not in observed_report_tables(result)[1][1]
    result["prep"]["data"]["Score"] = pd.Series([0, pd.NA, 2, 2, 2, 3], dtype="Int64")
    assert observed_report_tables(result) == []


def test_shared_person_pairs_distinguish_equal_marginal_workloads(writing_run):
    result, _ = writing_run
    summaries = []
    for pairs in (list(combinations("ABCD", 2)), [("A", "B")]*3 + [("C", "D")]*3):
        data = pd.DataFrame([(f"P{i}", rater, i % 4) for i, pair in enumerate(pairs)
                             for rater in pair], columns=["Person", "Rater", "Score"])
        result["prep"]["data"] = data
        summaries.append(observed_report_tables(result)[0][1].set_index("Characteristic")["Value"])
    for key in ("Retained ratings", "Rater levels", "Ratings per Rater level", "Unique Rater levels per person"):
        assert summaries[0][key] == summaries[1][key]
    assert summaries[0]["Rater pairs sharing persons"] == "6"
    assert summaries[1]["Rater pairs sharing persons"] == "2"
    assert summaries[0]["Shared persons per overlapping Rater pair"] == "1–1"
    assert summaries[1]["Shared persons per overlapping Rater pair"] == "3–3"


def test_pair_overlap_expansion_is_bounded(writing_run):
    result, _ = writing_run
    result["prep"]["data"] = pd.DataFrame({"Person": ["P1"]*101,
                                           "Rater": [f"R{i}" for i in range(101)], "Score": [1]*101})
    _, design, note = observed_report_tables(result)[0]
    assert not design.Characteristic.str.contains("pairs sharing persons").any()
    assert "omitted above 100 rater levels" in note


@pytest.mark.parametrize("success", [False, None, np.nan])
def test_unsuccessful_or_unknown_termination_masks_publication_uncertainty(writing_run, success):
    result, diagnostics = writing_run
    diagnostics["measures"]["SE_Method"] = "conditional information approximation"
    diagnostics["reliability"] = app.calc_reliability(diagnostics["measures"])
    assert app._publication_reliability_table(result, diagnostics)[0].Use.eq("Diagnostic").all()
    result["opt"].success = success
    diagnostics["measures"]["CI_Lower"] = [-.8, 0.]
    diagnostics["measures"]["CI_Upper"] = [0., .8]
    table, note = app._publication_measure_table(result, diagnostics)
    assert table["Basis"].eq("Withheld").all()
    assert table[["SE / SD", "Lower", "Upper"]].eq("").all().all()
    assert "SE/SD and interval cells are blank" in note
    assert any("Interpretation is withheld" in text for text in app._publication_report_intro(result, include_details=False))
    reliability, _ = app._publication_reliability_table(result, diagnostics)
    assert reliability.Use.eq("Withheld").all()
    assert reliability[["RMSE", "Separation", "Strata", "Reliability"]].eq("").all().all()
    assert diagnostics["measures"].SE.eq(.2).all()


def test_publication_formats_share_new_observed_tables_and_original_values(writing_run, monkeypatch):
    result, diagnostics = writing_run
    monkeypatch.setattr(app, "_publication_figure_payloads", lambda *a: [])
    tables = app._publication_report_tables(result, diagnostics, {}, include_details=False)
    assert [title for title, _, _ in tables] == [
        "Analysis Design and Retained Ratings", "Facet Estimates and Fit", "Observed Category Use"]
    html = app.build_publication_html_bytes(result, diagnostics, include_details=False).decode()
    word = Document(BytesIO(app.build_publication_word_bytes(result, diagnostics, include_details=False)))
    parsed = pd.read_html(StringIO(html))
    assert len(parsed) == len(word.tables) == 3
    assert parsed[0].Value.iloc[0] == "6"
    np.testing.assert_array_equal(parsed[2].Count, [2, 0, 3, 1])
    assert "Ratings: 6; persons: 2" in html
    assert "Ratings: 999" not in html
    assert "All recorded score categories are listed, including zero-count categories when present." in html
    assert word.tables[2].cell(2, 1).text == "0"


def test_actual_settings_and_descriptive_results_match_word_and_markdown(writing_run):
    result, diagnostics = writing_run
    context = app._manuscript_run_context(result, diagnostics)
    markdown = manuscript_markdown(result, run_context=context)
    word = Document(BytesIO(manuscript_word_bytes(result, run_context=context)))
    text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    for phrase in ("6 ratings and 2 person identifiers", "maximum iterations = 200",
                   "stopping tolerance = 1e-07", "Rater: +1", "Sum-to-zero constraints",
                   "12", "3 of 4 recorded score categories", "For Rater, 2 finite estimates",
                   "-0.400 to 0.400 logits", "not tests of differences"):
        assert phrase in markdown and phrase in text
    assert "Convergence status: converged" not in markdown
    assert "does not establish stationarity" in markdown
    assert "[Answer RQ1" in markdown


def test_appendix_and_results_draft_recompute_stale_counts_and_unused_categories(writing_run):
    result, diagnostics = writing_run
    result["prep"]["unused_score_categories"] = []
    for builder in (app.generate_method_appendix_text, app.generate_manuscript_reporting_template):
        text = builder(result, diagnostics)
        assert "unused categories retained: 1" in text
        assert "unused categories retained: none" not in text
        assert "999" not in text and "persons: 99" not in text
        result["prep"]["rating_max"] = None
        assert "Recorded score support requires review" in builder(result, diagnostics)
        result["prep"]["rating_max"] = 3
    handoff = app.generate_manuscript_handoff_markdown(result, diagnostics)
    assert "6 observations; 2 persons" in handoff
    checklist = app.build_manuscript_handoff_checklist(result, diagnostics)
    assert "6 observations and 2 persons" in checklist.loc[checklist.Step.eq(5), "WhyItMatters"].iloc[0]
    assert "does not establish stationarity" in checklist.loc[checklist.Step.eq(8), "WhyItMatters"].iloc[0]


@pytest.mark.parametrize("method,free_sd,boundary,identification", [
    ("MML", False, False, None), ("MML", True, False, None),
    ("JMLE", False, True, None), ("JMLE", False, False, "eta_rank_deficient"),
    ("JMLE", False, False, "audit_size_limit"),
])
def test_saved_ready_flags_and_stale_prose_cannot_bypass_direct_or_ui_manuscript_holds(
    writing_run, method, free_sd, boundary, identification,
):
    result, diagnostics = writing_run
    result["config"].update(method=method, estimate_population_sd=free_sd,
                            jmle_facet_boundary_review=boundary, mml_engine="joint",
                            population_prior_sd=1.1, quad_points=31)
    if identification:
        result["identifiability"] = {"status": identification}
    before = deepcopy(diagnostics)
    context = app._manuscript_run_context(result, diagnostics)
    assert context["hold"]
    assert not any("For Rater" in text for text in context["results"])
    stale = {"results": ["STALE APPROVED RANKING"], "hold": ""}
    for supplied in (None, context, stale):
        markdown = manuscript_markdown(result, run_context=supplied)
        word = Document(BytesIO(manuscript_word_bytes(result, run_context=supplied)))
        text = "\n".join(paragraph.text for paragraph in word.paragraphs)
        assert "Interpretation is withheld" in markdown and "Interpretation is withheld" in text
        assert "STALE APPROVED RANKING" not in markdown + text
        assert "6 ratings from 2 person identifiers" in markdown
        if method == "MML":
            assert "MML engine = joint" in markdown and "31 nodes" in markdown
    pd.testing.assert_frame_equal(diagnostics["measures"], before["measures"])


def test_optimizer_boolean_is_reported_as_termination_not_mathematical_convergence():
    text = optimizer_termination_text({"opt": SimpleNamespace(success=True, nit=8)})
    assert "stopping rule was reached" in text and "stationarity" in text
    assert "converged" not in text
    assert "not recorded" in optimizer_termination_text({"opt": SimpleNamespace(success=np.nan)})


@pytest.mark.parametrize('include_details', [False, True])
@pytest.mark.parametrize('kind', ['pdf', 'word', 'html'])
def test_reports_retain_first_read_cautions_and_unassessed_scope(writing_run, monkeypatch, kind, include_details):
    from xml.etree import ElementTree as ET
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    result, diagnostics = writing_run
    diagnostics.update(pca_enabled=True, pca={
        'eigenvalues': np.array([3.74, 1.2]),
        'stability_table': pd.DataFrame([{'PCAStabilityStatus': 'Review',
                                        'SensitivityFlagSummary': 'loading shift < 0.4 & overlap review'}]),
    })
    original = deepcopy(diagnostics)
    monkeypatch.setattr(app, '_publication_figure_payloads', lambda *a: [])
    # Check basic/detailed output wiring without rebuilding unrelated appendices.
    monkeypatch.setattr(app, 'generate_method_appendix_text', lambda *a: 'Method appendix')
    monkeypatch.setattr(app, 'generate_manuscript_reporting_template', lambda *a: 'Writing appendix')
    pdf_text = []
    original_build = SimpleDocTemplate.build
    def capture(self, story, *a, **kw):
        pdf_text.extend(part.getPlainText() for part in story if isinstance(part, Paragraph))
        return original_build(self, story, *a, **kw)
    monkeypatch.setattr(SimpleDocTemplate, 'build', capture)
    payload = getattr(app, f'build_publication_{kind}_bytes')(
        result, diagnostics, include_details=include_details)
    if kind == 'word':
        text = '\n'.join(p.text for p in Document(BytesIO(payload)).paragraphs)
    elif kind == 'html':
        text = '\n'.join(ET.fromstring(payload.decode().split('<!DOCTYPE html>\n', 1)[1]).itertext())
    else:
        assert payload.startswith(b'%PDF')
        text = '\n'.join(pdf_text)
    notes = text.split('Computation and interpretation checks', 1)[1].split('Results tables', 1)[0]
    for phrase in ('Optimizer reached its stopping rule.', 'does not establish model validity',
                   '3.74', 'Stability audit: Review', 'loading shift < 0.4 & overlap review',
                   'No bias/interaction table was computed', 'not mandatory for every study'):
        assert phrase in notes
    assert text.index('Computation and interpretation checks') < text.index('Results tables')
    pd.testing.assert_frame_equal(diagnostics['measures'], original['measures'])
    pd.testing.assert_frame_equal(diagnostics['pca']['stability_table'], original['pca']['stability_table'])
    np.testing.assert_array_equal(diagnostics['pca']['eigenvalues'], original['pca']['eigenvalues'])
    assert result['opt'].success is True


@pytest.mark.parametrize('case', ['failed', 'missing', 'truthy', 'fixed_sd', 'free_sd', 'boundary'])
def test_report_notes_keep_calibration_holds_without_reusing_stale_success(writing_run, case):
    result, diagnostics = writing_run
    result['summary'] = pd.DataFrame([{'Converged': True, 'InferenceReady': True}])
    if case == 'missing':
        result.pop('opt')
    elif case in {'failed', 'truthy'}:
        result['opt'].success = False if case == 'failed' else 'True'
    elif case in {'fixed_sd', 'free_sd'}:
        result['config'].update(method='MML', estimate_population_sd=case == 'free_sd')
    else:
        result['config']['jmle_facet_boundary_review'] = True
    notes = app._publication_review_notes(result, diagnostics)
    assert '1. Convergence (Do not interpret yet)' in notes
    assert '4. Dimensionality (Skipped)' in notes
    assert 'No bias/interaction table was computed' in notes
    assert '1. Convergence (OK)' not in notes
    assert result['summary'].InferenceReady.iloc[0]
