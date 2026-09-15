"""A citation follows its recorded software version, never the viewer's version."""

from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

import streamlit_app as app
from mfrm_app import citation, exports
from mfrm_app.manuscript import manuscript_markdown, manuscript_word_bytes


@pytest.mark.parametrize("version", [citation.METADATA["version"], "0.2.16-beta", None, ""])
def test_citation_follows_the_fit_into_manuscript_and_archives(version):
    result = {"config": {"app_version": version, "method": "MML", "model": "PCM",
                         "estimate_population_sd": True}, "raw_data": "PRIVATE_RATING_ROWS"}
    matched = version == citation.METADATA["version"]
    assets = app.build_evidence_contract_text_assets(result, None)
    markdown = manuscript_markdown(result)
    word = Document(BytesIO(manuscript_word_bytes(result)))
    word_text = "\n".join(p.text for p in word.paragraphs)
    assert (citation.APA_MARKDOWN in markdown) == matched
    assert (citation.APA in word_text) == matched
    assert "inference hold" in markdown
    assert "PRIVATE_RATING_ROWS" not in "".join(assets.values()) + markdown + word_text
    assert result["config"]["app_version"] == version
    assert result["config"]["estimate_population_sd"] is True
    binder = app.build_manuscript_binder_assets({}, assets)
    for bundle in (exports.build_tables_zip({}, assets), exports.build_mixed_asset_zip(binder)):
        with ZipFile(BytesIO(bundle)) as zipped:
            assert "software_citation.md" in zipped.namelist()
            assert ("CITATION.cff" in zipped.namelist()) == matched
            assert ("mfrm_software.bib" in zipped.namelist()) == matched
    if matched:
        reference = next(p for p in word.paragraphs if p.text == citation.APA)
        assert reference.paragraph_format.first_line_indent.inches == -0.5
        assert any(r.text == citation.METADATA["title"] and r.italic for r in reference.runs)
    else:
        assert citation.UNAVAILABLE in assets["software_citation.md"]
        assert citation.APA not in "".join(assets.values())


def test_bad_or_absent_result_metadata_never_generates_current_release_citation():
    for result in (None, {}, {"config": None}, {"config": []}, {"config": {"app_version": 2026}}):
        assert not citation.matches_result(result)
        assert set(citation.result_assets(result)) == {"software_citation.md"}


def test_publication_cites_software_only_when_both_version_and_narrative_match(monkeypatch):
    result = {"config": {"app_version": app.APP_VERSION, "method": "JMLE", "model": "RSM"}}
    monkeypatch.setattr(app, "_publication_figure_payloads", lambda *_: [])
    monkeypatch.setattr(app, "generate_method_appendix_text",
                        lambda r, *_: "# Methods\n\n" + citation.analysis_statement(r))
    monkeypatch.setattr(app, "generate_manuscript_reporting_template", lambda *_: "# Results\n\nReview required.")
    html = app.build_publication_html_bytes(result, {}).decode()
    word = Document(BytesIO(app.build_publication_word_bytes(result, {})))
    assert f"<em>{citation.METADATA['title']}</em>" in html
    assert citation.IN_TEXT in html
    refs = [p for p in word.paragraphs if p.text == citation.APA]
    assert len(refs) == 1
    assert any(r.text == citation.METADATA["title"] and r.italic for r in refs[0].runs)
    assert citation.APA not in app.build_apa_reference_list("No software citation here", software_result=result)
    old = {"config": {**result["config"], "app_version": "0.2.16-beta"}}
    assert citation.APA not in app.build_apa_reference_list(citation.IN_TEXT, software_result=old)
    assert citation.IN_TEXT not in app.build_publication_html_bytes(old, {}).decode()


def test_method_appendix_does_not_fill_missing_version_from_running_app():
    result = {"config": {"model": "RSM", "method": "JMLE"}, "opt": None}
    text = app.generate_method_appendix_text(result, {})
    assert "app version: not recorded" in text
    assert citation.IN_TEXT not in text
    assert app.APP_VERSION not in text
    result["config"]["app_version"] = app.APP_VERSION
    assert citation.IN_TEXT in app.generate_method_appendix_text(result, {})


@pytest.mark.parametrize("lang", ["en", "ja"])
def test_result_citation_downloads_disappear_when_recorded_version_changes(lang):
    def view(lang):
        import streamlit as st
        import streamlit_app as app
        st.session_state["lang"] = lang
        st.session_state.setdefault("result", {"config": {"app_version": app.APP_VERSION}})
        app.render_software_citation(key_prefix="result", result=st.session_state["result"])

    at = AppTest.from_function(view, args=(lang,)).run()
    assert not at.exception
    assert len(at.get("download_button")) == 2
    assert any(citation.APA == block.value for block in at.code)
    at.session_state["result"] = {"config": {"app_version": "0.2.16-beta"}}
    at.run()
    assert not at.exception
    assert not at.get("download_button") and not at.code
    assert len(at.info) == 1
    assert "software.result_unavailable" not in at.info[0].value
