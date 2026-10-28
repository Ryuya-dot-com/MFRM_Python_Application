"""An interval's own qualification survives reporting/replay; global holds do not lift."""
from copy import deepcopy
import json
import pickle

import numpy as np
import pandas as pd
import pytest

from mfrm_app import engine, output_qualification, structural_intervals as intervals


@pytest.fixture(scope='module', params=['RSM', 'PCM'])
def fit(request):
    data = pd.DataFrame([dict(Person=f'{p:03d}', Rater=f'R{r}', Task=f'T{t}', Score=(p+2*r+t)%3)
                         for p in range(30) for r in range(3) for t in range(2)])
    # The ordinary default origin is intentionally retained.
    return engine.mfrm_estimate(data, 'Person', ['Rater', 'Task'], 'Score',
        model=request.param, method='MML', step_facet='Task', quad_points=61, maxit=100)


def test_qualified_difference_uses_full_covariance_and_replays_without_refitting(fit, monkeypatch):
    original = pickle.dumps(fit)
    monkeypatch.setattr(engine, 'minimize', lambda *a, **k: pytest.fail('Interval calculation must not refit'))
    record = intervals.compute_structural_intervals(fit, facet='Rater', difference=['R0', 'R2'])
    table = intervals.read_structural_intervals(fit, json.loads(intervals.interval_json(record)))
    assert table.CIEligible.all() and not table.CoverageQualified.any()
    assert {row['points'] for row in record['calculations']} == {61, 121}
    assert fit['config']['noncenter_facet'] == 'Person'
    prepared = intervals.review.prepare_fixed_sd_check(fit)
    independent = engine._compute_mml_observed_information(prepared['coordinates'], prepared['idx'],
        prepared['config'], prepared['sizes'], engine.make_mml_quadrature(prepared['config']), rel_step=1e-4)
    np.testing.assert_allclose(record['calculations'][0]['information'], independent, rtol=1e-7, atol=2e-7)
    covariance = np.asarray(record['target_covariance'])
    difference = table.iloc[-1]
    assert difference.Estimate == pytest.approx(table.iloc[0].Estimate-table.iloc[2].Estimate)
    assert difference.SE**2 == pytest.approx(covariance[0,0]+covariance[2,2]-2*covariance[0,2])
    assert abs(difference.SE**2 - (covariance[0,0]+covariance[2,2])) > .01
    assert output_qualification.mml_inference_hold_reason(fit['config'])
    assert not output_qualification.facet_equivalence_qualification().public_conclusion_allowed
    assert pickle.dumps(fit) == original
    # After recording, likelihood evaluation is forbidden during read/HTML/figure.
    monkeypatch.setattr(engine, 'mfrm_loglik_mml_value_grad', lambda *a, **k: pytest.fail('Replay evaluated likelihood'))
    text = intervals.interval_html(fit, record)
    assert 'CIEligible' in text and 'CoverageQualified' in text and 'finite-sample' in text
    figure = intervals.interval_figure(table)
    np.testing.assert_allclose(figure.data[0].error_x.array, table.Upper-table.Estimate)


def test_full_covariance_gate_detects_unchanged_marginal_ses():
    a = np.array([[1,.4,-.4],[.4,1,0],[-.4,0,1]])
    b = np.array([[1,-.4,.4],[-.4,1,0],[.4,0,1]])
    j = np.vstack([np.eye(3), -np.ones(3)])
    np.testing.assert_allclose(np.diag(j@a@j.T), np.diag(j@b@j.T))
    record = dict(coordinates=[0.,0.,0.], optimizer_success=True,
        boundary=dict(structural_design_rank=3,status='no_structural_ray_detected',no_finite_mle_certified=False),
        calculations=[dict(points=q, information=np.linalg.inv(v).tolist(), gradient=[0.,0.,0.])
                      for q,v in [(31,a),(61,b)]])
    checks, _ = intervals._assess(record)
    check = next(row for row in checks if row['check']=='full_covariance_change')
    assert not check['passed'] and check['value'] > 1


@pytest.mark.parametrize('change', ['data','point','success','record','eligibility'])
def test_stale_or_edited_records_cannot_restore_qualification(fit, change):
    record = intervals.compute_structural_intervals(fit, facet='Rater')
    changed = deepcopy(fit)
    if change=='data': changed['prep']['data'].iloc[0, changed['prep']['data'].columns.get_loc('Score')] = 2
    elif change=='point': changed['opt'].x[0] += .1
    elif change=='success': changed['opt'].success = False
    elif change=='record': record['calculations'][0]['gradient'][0] += 10
    else:
        record['table'][0]['CIEligible'] = False
        record['record_hash'] = intervals._fingerprint(record)
    with pytest.raises(ValueError): intervals.read_structural_intervals(changed, record)


def test_termination_failure_and_information_failure_withhold_bounds(fit):
    changed = deepcopy(fit); changed['opt'].success = False
    record = intervals.compute_structural_intervals(changed, facet='Rater')
    table = intervals.read_structural_intervals(changed, record)
    assert not table.CIEligible.any() and table.SE.isna().all() and table.Lower.isna().all()
    assert table.Reason.str.contains('optimizer_termination').all()
    assert np.isfinite(table.Estimate).all()
    record['calculations'][0]['information'] = np.zeros_like(record['calculations'][0]['information']).tolist()
    rows, checks, cov = intervals._table(record)
    assert cov is None and all(row['SE'] is None for row in rows)
    assert not next(c for c in checks if c['check']=='GH61:positive_information')['passed']


def test_report_formats_use_the_selected_interval_qualification(fit):
    import streamlit_app as app
    changed = deepcopy(fit)
    changed['_structural_intervals'] = intervals.compute_structural_intervals(changed, facet='Rater', difference=['R0','R1'])
    tables = app._publication_report_tables(changed, {'measures':pd.DataFrame()}, {}, include_details=False)
    _, table, note = next(row for row in tables if row[0]=='Specified Structural Intervals')
    assert table.Use.eq('Approximate').all() and table.Interval.str.startswith('[').all()
    assert 'pointwise' in note.lower() and 'coverage' in note
    from io import BytesIO
    from docx import Document
    html = app.build_publication_html_bytes(changed, {'measures':pd.DataFrame()}, include_details=False).decode()
    word = Document(BytesIO(app.build_publication_word_bytes(changed, {'measures':pd.DataFrame()}, include_details=False)))
    word_text = '\n'.join(p.text for p in word.paragraphs)+'\n'+'\n'.join(
        cell.text for t in word.tables for row in t.rows for cell in row.cells)
    for text in (html, word_text):
        assert 'Specified Structural Intervals' in text and 'Approximate' in text
        assert 'coverage guarantees' in text
    assert app.build_publication_pdf_bytes(changed, {'measures':pd.DataFrame()}, include_details=False).startswith(b'%PDF-')
    changed['opt'].fun += .1
    tables = app._publication_report_tables(changed, {'measures':pd.DataFrame()}, {}, include_details=False)
    _, table, _ = next(row for row in tables if row[0]=='Specified Structural Intervals')
    assert 'Unavailable' in table.Status.iloc[0]


def test_medium_sized_input_is_limited_by_arrays_not_the_old_row_cap(fit):
    changed = deepcopy(fit)
    data = changed['prep']['data']
    pieces = [data.assign(Person=data.Person.astype(str)+f'_{i}') for i in range(17)]
    prep = engine.prepare_mfrm_data(pd.concat(pieces), 'Person', ['Rater','Task'], 'Score')
    changed['prep'] = prep
    changed['config']['n_person'] = len(prep['levels']['Person'])
    assert len(prep['data']) > 3000
    prepared = intervals.review.prepare_fixed_sd_check(changed)
    assert prepared['problem'].observations == len(prep['data'])
